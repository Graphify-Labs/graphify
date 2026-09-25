"""Tests for graphify.detect.docx_to_markdown."""
from __future__ import annotations

import pytest

docx = pytest.importorskip("docx")
from docx.oxml import parse_xml  # noqa: E402
from docx.shared import Inches  # noqa: E402

from graphify.detect import docx_to_markdown  # noqa: E402

_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _append_xml(document, *fragments: str) -> None:
    section_properties = document.element.body[-1]
    body = parse_xml(f"<w:body {_NAMESPACES}>{''.join(fragments)}</w:body>")
    for element in list(body):
        section_properties.addprevious(element)


def _markdown(document, tmp_path) -> list[str]:
    path = tmp_path / "doc.docx"
    document.save(path)
    return docx_to_markdown(path).split("\n")


def test_plain_paragraphs_keep_their_markdown(tmp_path):
    document = docx.Document()
    document.add_heading("Title", level=1)
    document.add_heading("Section", level=2)
    document.add_paragraph("Body text.")
    document.add_paragraph("")
    document.add_paragraph("An item", style="List Bullet")
    # A tab stop is a w:tab in the paragraph properties, not text.
    document.add_paragraph("Tab stop").paragraph_format.tab_stops.add_tab_stop(Inches(1))

    assert _markdown(document, tmp_path) == [
        "# Title",
        "## Section",
        "Body text.",
        "",
        "- An item",
        "Tab stop",
    ]


def test_a_table_stays_under_its_heading(tmp_path):
    document = docx.Document()
    document.add_heading("Pricing", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Basic"
    table.cell(1, 1).text = "10 EUR"
    document.add_heading("Notes", level=1)

    assert _markdown(document, tmp_path) == [
        "# Pricing",
        "| Plan | Price |",
        "| --- | --- |",
        "| Basic | 10 EUR |",
        "# Notes",
    ]


def test_a_cell_stays_on_its_row(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Basic | Pro"
    table.cell(1, 1).text = "10 EUR\nper month"  # a line break inside the paragraph
    table.cell(1, 1).add_paragraph("billed yearly")

    assert _markdown(document, tmp_path)[-1] == "| Basic \\| Pro | 10 EUR per month billed yearly |"


def test_a_nested_table_is_flattened_into_its_cell_once(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Outer"
    nested = table.cell(0, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Inner A"
    nested.cell(0, 1).text = "Inner B"

    assert _markdown(document, tmp_path)[0] == "| Outer | Inner A Inner B |"


def test_tracked_changes_content_controls_and_text_boxes(tmp_path):
    box = "<w:txbxContent><w:p>" + _run("Callout") + "</w:p></w:txbxContent>"
    document = docx.Document()
    _append_xml(
        document,
        "<w:p>"
        + _run("The fee is ")
        + '<w:del w:id="1" w:author="a"><w:r><w:delText>ten</w:delText><w:tab/></w:r></w:del>'
        + f'<w:ins w:id="2" w:author="a">{_run("twelve")}</w:ins>'
        + _run(" euros.")
        + "</w:p>",
        "<w:p>" + _run("Client: ") + f"<w:sdt><w:sdtPr/><w:sdtContent>{_run('Acme')}</w:sdtContent></w:sdt></w:p>",
        f"<w:sdt><w:sdtPr/><w:sdtContent><w:p>{_run('Block control')}</w:p></w:sdtContent></w:sdt>",
        # Word writes every text box twice: a DrawingML shape and a VML fallback copy.
        "<w:p>"
        + _run("Host")
        + "<w:r><mc:AlternateContent>"
        + f'<mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp></w:drawing></mc:Choice>'
        + f"<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>"
        + "</mc:AlternateContent></w:r></w:p>",
    )

    assert _markdown(document, tmp_path) == [
        "The fee is twelve euros.",
        "Client: Acme",
        "Block control",
        "Host",
        "Callout",
    ]
